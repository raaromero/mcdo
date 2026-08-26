"""Enforce docs/STANDARDS.md so presentation defects cannot regress.

These caught real problems repeatedly: underscored display strings leaking into
figures (``X_NA``, ``alpha_t``), internal planning vocabulary in figure titles
("Phase 3"), and deck entries pointing at figures nothing produces.
"""

import os
import re
import sys

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
SCRIPTS = os.path.join(ROOT, "scripts")

# Calls whose string argument is rendered for a reader.
DISPLAY_CALL = re.compile(
    r"(?:set_title|set_xlabel|set_ylabel|suptitle)\(\s*[a-z]?(['\"])(.*?)\1", re.S
)
# Underscored identifiers that must never appear in rendered text.
BAD_TOKEN = re.compile(r"(?<![\\$])\b(?:X_NA|alpha_t|N_theta|N_freq|lam_c|r_ap|w_beam)\b|α_t")
PHASE = re.compile(r"\bPhase\s+\d", re.I)


def _script_files():
    return sorted(
        os.path.join(SCRIPTS, f)
        for f in os.listdir(SCRIPTS)
        if f.endswith(".py") and not f.startswith("_")
    )


def _display_strings(path):
    with open(path, encoding="utf-8") as fh:
        src = fh.read()
    return [m.group(2) for m in DISPLAY_CALL.finditer(src)]


@pytest.mark.parametrize("path", _script_files(), ids=os.path.basename)
def test_no_underscored_identifiers_in_rendered_text(path):
    """Axis labels and titles must not expose code identifiers."""
    offenders = [s for s in _display_strings(path) if BAD_TOKEN.search(s)]
    assert not offenders, (
        f"{os.path.basename(path)} renders code identifiers: {offenders}. "
        "Use reader-facing wording (see docs/STANDARDS.md)."
    )


@pytest.mark.parametrize("path", _script_files(), ids=os.path.basename)
def test_no_internal_phase_vocabulary_in_titles(path):
    """'Phase N' is planning vocabulary, not physics; it must not reach a figure."""
    offenders = [s for s in _display_strings(path) if PHASE.search(s)]
    assert not offenders, (
        f"{os.path.basename(path)} puts internal phase labels in a figure: {offenders}"
    )


def test_deck_manifest_is_consistent():
    """Every declared figure is used, every used figure is declared."""
    from mcdo.deck import DECK, FIGURES

    used = {item[3] for item in DECK if item[0] == "slide" and item[3]}
    undeclared = used - set(FIGURES)
    assert not undeclared, f"deck slides reference undeclared figures: {sorted(undeclared)}"

    unused = set(FIGURES) - used
    assert not unused, f"FIGURES declares figures no slide uses: {sorted(unused)}"


def test_deck_figures_name_a_real_generating_script():
    """A missing figure must be recoverable: the named script has to exist."""
    from mcdo.deck import FIGURES

    missing = [
        (key, script)
        for key, (_rel, script) in FIGURES.items()
        if not os.path.exists(os.path.join(ROOT, script))
    ]
    assert not missing, f"figures name scripts that do not exist: {missing}"


def test_deck_latex_keys_resolve():
    """Slides may only reference equations that exist."""
    from mcdo.deck import DECK, LX

    referenced = {item[4] for item in DECK if item[0] == "slide" and item[4]}
    unknown = referenced - set(LX)
    assert not unknown, f"slides reference unknown equations: {sorted(unknown)}"


def _estimated_text_height(shape):
    """Rough rendered height of a text box, in inches.

    Approximates the wrapped line count from the character count and the box
    width, using an average glyph advance of about 0.5 em for Helvetica. It is
    deliberately conservative: it catches bullets that clearly overrun their
    box, not borderline cases.
    """
    width_in = shape.width / 914400
    total = 0.0
    for para in shape.text_frame.paragraphs:
        text = "".join(r.text for r in para.runs)
        if not text:
            continue
        size_pt = max((r.font.size.pt for r in para.runs if r.font.size), default=18)
        indent_in = 0.35 * para.level
        chars_per_line = max(8, int((width_in - indent_in) * 72 / (size_pt * 0.5)))
        lines = max(1, -(-len(text) // chars_per_line))
        total += lines * size_pt * 1.22 / 72 + 9 / 72       # line height + paragraph space
    return total


def test_slide_text_fits_its_box():
    """Bullet text must not obviously overrun the box it was given."""
    from pptx import Presentation

    deck = os.path.join(ROOT, "output", "deck", "mcdo_results_deck.pptx")
    if not os.path.exists(deck):
        pytest.skip("deck not built yet")

    offenders = []
    for index, slide in enumerate(Presentation(deck).slides, start=1):
        for shape in slide.shapes:
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            if shape.top <= 500000:            # the title band, single line
                continue
            box_in = shape.height / 914400
            needed = _estimated_text_height(shape)
            if needed > box_in * 1.15:         # 15 per cent slack for the estimate
                first = shape.text_frame.text.strip().split("\n")[0][:44]
                offenders.append(f"slide {index}: needs ~{needed:.1f} in of {box_in:.1f} in — {first}")
    assert not offenders, "text overruns its box on:\n  " + "\n  ".join(offenders[:12])


def _deck_scripts():
    from mcdo.deck import FIGURES

    return sorted({os.path.join(ROOT, script) for _rel, script in FIGURES.values()})


def _titles(path):
    pattern = re.compile(r"(?:set_title|suptitle)\(\s*f?(['\"])(.*?)\1", re.S)
    with open(path, encoding="utf-8") as fh:
        return [m.group(2) for m in pattern.finditer(fh.read())]


@pytest.mark.parametrize("path", _deck_scripts(), ids=os.path.basename)
def test_figure_titles_are_plain_and_descriptive(path):
    """Deck figure titles: descriptive, no asides, no line breaks, sentence case."""
    offenders = []
    for raw in _titles(path):
        # f-string placeholders are code, not rendered text
        title = re.sub(r"\{[^{}]*\}", "", raw)
        core = re.sub(r"^\([a-zA-Z0-9]\)\s*", "", title).strip()
        if "[" in title:
            offenders.append(f"bracketed aside: {title[:60]}")
        elif "\\n" in title:
            offenders.append(f"line break: {title[:60]}")
        elif re.search(r"[A-Za-z]_[A-Za-z]", title):
            offenders.append(f"underscore: {title[:60]}")
        elif core[:1].islower():
            offenders.append(f"lowercase start: {title[:60]}")
    assert not offenders, f"{os.path.basename(path)}:\n  " + "\n  ".join(offenders)


@pytest.mark.parametrize("path", _deck_scripts(), ids=os.path.basename)
def test_figures_are_high_resolution(path):
    """Every saved figure must be at least 300 dpi."""
    with open(path, encoding="utf-8") as fh:
        src = fh.read()
    low = [d for d in re.findall(r"dpi\s*[=:]\s*(\d+)", src) if int(d) < 300]
    assert not low, f"{os.path.basename(path)} saves at dpi {low}, below the 300 standard"


CUT_WORD = re.compile(r"[xy][- ]cut", re.I)


@pytest.mark.parametrize("path", _deck_scripts(), ids=os.path.basename)
def test_no_cut_jargon_in_rendered_text(path):
    """'x-cut'/'y-cut' is internal jargon; rendered text names the component."""
    pattern = re.compile(
        r"(?:label\s*=|set_title|set_xlabel|set_ylabel|suptitle)\(?\s*f?(['\"])(.*?)\1", re.S
    )
    with open(path, encoding="utf-8") as fh:
        src = fh.read()
    offenders = [m.group(2)[:60] for m in pattern.finditer(src) if CUT_WORD.search(m.group(2))]
    assert not offenders, f"{os.path.basename(path)} renders cut jargon: {offenders}"


@pytest.mark.parametrize("path", _deck_scripts(), ids=os.path.basename)
def test_plot_line_widths_meet_standard(path):
    """Data curves are at least 1.6 pt so they survive projection."""
    offenders = []
    with open(path, encoding="utf-8") as fh:
        for number, line in enumerate(fh, start=1):
            if re.search(r"\.(?:plot|semilogy|loglog|semilogx)\(", line) \
                    and "axhline" not in line and "axvline" not in line:
                for w in re.findall(r"lw=([\d.]+)", line):
                    if 0 < float(w) < 1.6:   # lw=0 is a marker-only series
                        offenders.append(f"line {number}: lw={w}")
    assert not offenders, f"{os.path.basename(path)} has thin data lines: {offenders}"



def test_declared_figures_exist_on_disk():
    """A slide must never be built around a figure that is not there."""
    from mcdo.deck import FIGURES, OUT
    missing = [rel for rel, _ in FIGURES.values()
               if not os.path.exists(os.path.join(OUT, rel))]
    assert not missing, f"declared but absent: {sorted(missing)[:8]}"


def test_panel_index_matches_panel_content():
    """A slide's figure must be the panel it means, not whatever took that index.

    Every panel records its own title. Where a panel title carries a numeric
    setting (a numerical aperture, an obstruction ratio, a pulse duration), the
    slide that uses it has to name the same value.
    """
    import json
    from mcdo.deck import DECK, FIGURES, OUT

    titles = {}
    for root, _dirs, files in os.walk(OUT):
        for name in files:
            if name.endswith("_panels.json"):
                with open(os.path.join(root, name)) as fh:
                    for panel, title in json.load(fh).items():
                        rel = os.path.relpath(os.path.join(root, panel), OUT)
                        titles[rel] = title

    number = re.compile(r"\d+\.\d+")
    offenders = []
    for item in DECK:
        if item[0] != "slide" or item[3] not in FIGURES:
            continue
        rel = FIGURES[item[3]][0]
        panel_title = titles.get(rel)
        if not panel_title:
            continue
        wanted = set(number.findall(panel_title))
        if not wanted:
            continue
        text = item[1] + " " + " ".join(b[1] for b in item[2])
        if not wanted & set(number.findall(text)):
            offenders.append(f"{item[3]}: slide {item[1]!r} vs panel {panel_title!r}")
    assert not offenders, "slide and panel disagree:\n  " + "\n  ".join(offenders[:10])


@pytest.mark.parametrize("path", _deck_scripts(), ids=os.path.basename)
def test_no_figure_numbers_in_titles(path):
    """A plot title names what is shown, never 'Fig. N' from a source paper."""
    pattern = re.compile(
        r"(?:set_title|suptitle)\(\s*r?f?(['\"])(.*?)\1", re.S)
    with open(path, encoding="utf-8") as fh:
        src = fh.read()
    bad = re.compile(r"\b(fig|figure)\.?\s*\d", re.I)
    offenders = [m.group(2)[:60] for m in pattern.finditer(src) if bad.search(m.group(2))]
    assert not offenders, f"{os.path.basename(path)} names a figure number: {offenders}"
